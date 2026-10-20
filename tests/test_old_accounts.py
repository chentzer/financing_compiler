"""Accounts kept by hand before fincomp (template layout) copied in."""

import io
from datetime import datetime

import openpyxl

from fincomp import webapi
from fincomp import settings as S
from fincomp.config import Config
from fincomp.old_accounts import is_workbook, rows_from_workbook


def workbook() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Format"  # the template's blank tab: skipped
    ws = wb.create_sheet("October 2022")
    ws["A1"], ws["N1"] = "Spendings", "Earnings"
    for col, h in zip("ABCDEF", ["Date", "Item", "Category", "Amount (RM)", "Importance", "Remarks"]):
        ws[f"{col}2"] = h
    for col, h in zip("NOPQ", ["Date", "Item", "Amount (RM)", "Remarks"]):
        ws[f"{col}2"] = h
    ws["H2"], ws["I2"] = "Row Labels", "Sum of Amount (RM)"  # an old pivot table: ignored
    ws.append([datetime(2022, 10, 1), "Mouse", "Shopping", 27.97, "Nice to Have", "Lazada"])
    ws.append([datetime(2022, 10, 1), "Mouse", "Shopping", 27.97, "Nice to Have", ""])  # same again: both kept
    ws.append([datetime(2022, 10, 5), "Nasi lemak", "Food", 6, "Essential"])
    for r, row in enumerate([[datetime(2022, 10, 12), "Lazada refund", 34.79, "Wifi card"]], start=3):
        for col, v in zip("NOPQ", row):
            ws[f"{col}{r}"] = v
    ws["H3"], ws["I3"] = "Shopping", 55.94
    wb.create_sheet("Year End Summary")["A1"] = "Total Earnings"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_rows_copied_as_they_are():
    data = workbook()
    assert is_workbook(data) and not is_workbook(b"%PDF-1.7")
    (title, rows), = rows_from_workbook(data).items()
    assert title == "October 2022"
    assert [(r["kind"], r["date"], r["item"], r["amount"], r.get("category"), r["remarks"]) for r in rows] == [
        ("spend", "2022-10-01", "Mouse", 27.97, "Shopping", "Lazada"),
        ("spend", "2022-10-01", "Mouse", 27.97, "Shopping", ""),
        ("spend", "2022-10-05", "Nasi lemak", 6.0, "Food", ""),
        ("earn", "2022-10-12", "Lazada refund", 34.79, None, "Wifi card"),
    ]
    assert len({r["key"] for r in rows}) == 4 and all(r["key"].startswith("m-") for r in rows)
    assert rows_from_workbook(data)["October 2022"][0]["key"] == rows[0]["key"]  # same every time


def test_into_a_new_year_sheet_and_not_twice():
    tabs = S.tabs_from_config(Config())
    tabs[S.REMOVED] = [S.HEADERS[S.REMOVED]]
    webapi.use_settings(tabs)
    out = webapi.old_accounts([{"name": "2022 Accounts.xlsx", "data": workbook()}])
    assert out["years"] == {"October 2022": 2022}
    rows = out["months"]["October 2022"]

    titles = ["Sheet1"]
    prep = webapi.prepare_rows(2022, titles, rows, "October 2022", created=True)
    titles = ["To Update", "October 2022", "Year End Summary"]
    ids = {t: i for i, t in enumerate(titles)}
    done = webapi.finish_rows(2022, ids, {r: [] for r in prep["reads"]})
    data = {d["range"]: d["values"] for d in done["data"]}
    assert done["added"] == 4
    assert data["'October 2022'!A3:H5"][2][:7] == ["2022-10-05", "Nasi lemak", "Food", 6.0, "Essential", "", ""]
    assert data["'October 2022'!P3:U3"][0][:5] == ["2022-10-12", "Lazada refund", 34.79, "", "Wifi card"]

    # Importing again: the rows are there by key, nothing is added.
    prep = webapi.prepare_rows(2022, titles, rows, "October 2022")
    sheet = {
        "'October 2022'!A3:H": [[r["date"], r["item"], r["category"], r["amount"], "", "", "", r["key"]]
                                for r in rows if r["kind"] == "spend"],
        "'October 2022'!P3:U": [[r["date"], r["item"], r["amount"], "", "", r["key"]] for r in rows if r["kind"] == "earn"],
    }
    again = webapi.finish_rows(2022, ids, {r: sheet.get(r, []) for r in prep["reads"]})
    assert again["added"] == 0


def test_month_tabs_without_the_year_and_amount_only_rows():
    from fincomp.old_accounts import sheets_from_workbook

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "January"  # 2024's workbook names months without the year
    for col, h in zip("ABCDEF", ["Date", "Item", "Category", "Amount (RM)", "Importance", "Remarks"]):
        ws[f"{col}2"] = h
    ws.append([datetime(2024, 1, 3), "Lunch", "Food", 12.0, "Essential"])
    ny = wb.create_sheet("New Years")
    ny["A1"] = "Ang Pao"
    for col, h in zip("ABCD", ["Date", "Person", "Amount (RM)", "Remarks"]):
        ny[f"{col}2"] = h
    ny["C3"], ny["C4"] = 100, 50  # only the amounts were kept
    ny["A5"], ny["C5"] = "Total", "=SUM(C3:C4)"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    months = rows_from_workbook(data)
    assert list(months) == ["January 2024"] and months["January 2024"][0]["item"] == "Lunch"
    (sheet,) = sheets_from_workbook(data)
    assert sheet["name"] == "New Years" and sheet["year"] == 2024
    assert [r["amount"] for r in sheet["tables"][0]["rows"]] == [100.0, 50.0]  # the total row isn't a row
