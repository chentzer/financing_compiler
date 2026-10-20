from datetime import date
from pathlib import Path

import gspread
import openpyxl
import pytest

from fincomp import gsheets
from fincomp.accounts import build_rows, month_layout, month_links, new_rows, plan, trip_layout
from fincomp.accounts_xlsx import write_accounts_xlsx
from fincomp.config import Config, Trip
from fincomp.models import Transaction

RULES = [
    {"match": "GRAB", "category": "Transport", "importance": "Essential"},
    {"match": ["KK MART", "FAMILYMART"], "category": "Food"},
    {"match": "SPOTIFY", "category": "Entertainment", "importance": "Nice to Have", "item": "Spotify"},
    {"match": "IGNORE ME", "skip": True},
    {"match": "WEIRD", "category": "Not A Real Category"},
    {"match": "KLOOK", "trip": "Japan Trip"},
]

JAPAN = Trip(
    name="Japan Trip", start=date(2026, 11, 8), end=date(2026, 11, 20), currency="Yen",
    match=["AIRASIA", "AGODA"], link="https://example.com/plan",
)


def tx(d, amount, desc, source="TNG eWallet", internal=False):
    return Transaction(date=d, source=source, account="eWallet", description=desc, amount=amount, internal=internal)


def sample():
    return [
        tx(date(2026, 7, 1), 3000.0, "SALARY", source="CIMB"),
        tx(date(2026, 7, 2), -12.0, "GRAB* RIDE"),
        tx(date(2026, 7, 3), -5.0, "KK MART"),
        tx(date(2026, 7, 3), -5.0, "KK MART"),  # same coffee twice
        tx(date(2026, 7, 4), -15.9, "SPOTIFY P0123"),
        tx(date(2026, 7, 5), -50.0, "RELOAD", internal=True),
        tx(date(2026, 7, 6), -1.0, "IGNORE ME"),
        tx(date(2026, 7, 7), -9.0, "WEIRD SHOP"),
        tx(date(2026, 8, 1), -20.0, "UNKNOWN MERCHANT"),
    ]


def japan_sample():
    return [
        tx(date(2026, 8, 17), -225.0, "AIRASIA BERHAD", source="Maybank Card"),  # booked early
        tx(date(2026, 3, 1), -99.0, "AIRASIA BERHAD", source="Maybank Card"),  # too early: normal spending
        tx(date(2026, 10, 2), -80.0, "KLOOK TRAVEL", source="Maybank Card"),  # via a rule
        tx(date(2026, 11, 3), -790.67, "AGODA HOTEL OSAKA", source="Maybank Card"),
        tx(date(2026, 11, 14), -131.98, "DON QUIJOTE", source="Maybank Card"),  # during the trip
        tx(date(2026, 11, 15), 200.0, "FRIEND PAYBACK"),  # money in stays on the month tab
        tx(date(2026, 11, 25), -9.0, "KK MART"),  # after the trip
    ]


def config(**kw):
    return Config(rules=RULES, **kw)


def test_rows_follow_template_and_rules():
    rows = build_rows(sample(), config())
    july = [r for r in rows if r.date.month == 7]
    got = [(r.kind, r.item, r.amount, r.category, r.importance) for r in july]
    assert got == [
        ("earn", "SALARY", 3000.0, "", ""),
        ("spend", "GRAB* RIDE", 12.0, "Transport", "Essential"),
        ("spend", "KK MART", 5.0, "Food", ""),
        ("spend", "KK MART", 5.0, "Food", ""),
        ("spend", "Spotify", 15.9, "Entertainment", "Nice to Have"),
        ("spend", "WEIRD SHOP", 9.0, "", ""),  # not in the dropdown list -> blank
    ]
    # Internal and skipped rows are left out; repeats get distinct keys.
    assert len({r.key for r in rows}) == len(rows)


def test_keys_are_stable_across_runs():
    a = [r.key for r in build_rows(sample(), config())]
    assert a == [r.key for r in build_rows(sample(), config())]
    assert new_rows(build_rows(sample(), config()), set(a[:3])) == build_rows(sample(), config())[3:]


def test_trip_spending_goes_to_trip_tab():
    tabs = plan(build_rows(japan_sample(), config(trips=[JAPAN])), config(trips=[JAPAN]))[2026]
    assert [t.title for t in tabs] == ["March 2026", "November 2026", "Japan Trip"]
    by_title = {t.title: [(r.item, r.amount) for r in t.rows] for t in tabs}
    assert by_title["Japan Trip"] == [
        ("AIRASIA BERHAD", 225.0), ("KLOOK TRAVEL", 80.0), ("AGODA HOTEL OSAKA", 790.67), ("DON QUIJOTE", 131.98),
    ]
    assert by_title["November 2026"] == [("FRIEND PAYBACK", 200.0), ("KK MART", 9.0)]
    assert by_title["March 2026"] == [("AIRASIA BERHAD", 99.0)]


def test_trip_without_config_stays_on_month_tab():
    # The KLOOK rule names a trip that isn't configured.
    tabs = plan(build_rows(japan_sample(), config()), config())[2026]
    assert "Japan Trip" not in [t.title for t in tabs]
    assert sum(len(t.rows) for t in tabs) == 7


def test_month_layout_positions():
    cells = month_layout(Config(), open_ended=True)
    assert [cells[f"{c}2"] for c in "ABCDEFGH"] == [
        "Date", "Item", "Category", "Amount (RM)", "Importance", "Source", "Remarks", "Key"]
    assert [cells[f"{c}2"] for c in "PQRSTU"] == ["Date", "Item", "Amount (RM)", "Source", "Remarks", "Key"]
    # A blank column between sections.
    assert not any(f"{c}2" in cells for c in "ILOV")
    assert (cells["J2"], cells["K2"], cells["M2"], cells["W2"]) == ("Category", "Spent (RM)", "Importance", "Summary")
    assert cells["J3"] == "Bills" and cells["K3"] == '=SUMIF($C$3:$C,"Bills",$D$3:$D)'
    assert cells["J11"] == "Not set" and cells["J12"] == "Total"
    assert cells["M3"] == "Essential" and cells["M9"] == "Total"
    assert cells["X3"] == "=SUM($R$3:$R)" and cells["W5"] == "Trips" and cells["W6"] == "Final"


def test_trip_layout_matches_japan_sheet():
    cells = trip_layout(JAPAN, Config(), open_ended=True)
    assert cells["A1"] == "Japan Trip" and cells["C1"] == '=HYPERLINK("https://example.com/plan","Trip Link")'
    assert [cells[f"{c}3"] for c in "ABCDEFG"] == [
        "Date", "Item", "Category", "Amount (RM)", "Importance", "Source", "Remarks"]
    # Paid back to you (RM), like the 2025 trip sheets.
    assert [cells[f"{c}3"] for c in "JKLMN"] == ["Date", "From", "Amount (RM)", "Remarks", "Key"]
    # The yen spending block starts at S; money received in other currencies after it.
    assert cells["S2"] == "Spending in Yen" and cells["V3"] == "Amount (Yen)"
    assert [cells[f"{c}3"] for c in ("Z", "AA", "AB", "AC", "AD")] == ["Date", "Item", "Currency", "Amount", "Remarks"]
    # Money Exchange is paid in RM but spent as yen, so it isn't counted twice.
    assert cells["Q5"] == '=SUM($D$4:$D)-SUMIF($B$4:$B,"Money Exchange",$D$4:$D)'
    assert cells["Q6"] == "=SUM($L$4:$L)" and cells["Q8"] == "=SUM($V$4:$V)" and cells["Q10"] == "=Q8*Q9"
    assert cells["Q4"] == "=Q5-Q6+Q10"  # what you really paid
    domestic = trip_layout(Trip("Penang", date(2026, 1, 1), date(2026, 1, 3)), Config(), open_ended=True)
    assert domestic["Q4"] == "=Q5-Q6" and "S2" not in domestic


def test_trip_with_several_currencies():
    europe = Trip("Europe", date(2024, 10, 20), date(2024, 10, 24), currency="CHF", rate=5.1,
                  extra=[("EUR", 4.9), ("GBP", None)])
    cells = trip_layout(europe, Config(), open_ended=True)
    assert (cells["S2"], cells["Z2"], cells["AG2"]) == ("Spending in CHF", "Spending in EUR", "Spending in GBP")
    assert cells["AC3"] == "Amount (EUR)" and cells["AN3"] == "Date" and cells["AP3"] == "Currency"  # received after GBP
    assert (cells["Q9"], cells["Q13"], cells["Q17"]) == (5.1, 4.9, "")  # each currency's rate
    assert cells["Q4"] == "=Q5-Q6+Q10+Q14+Q18"


def test_final_subtracts_trip_in_its_end_month():
    links = month_links(["October 2026", "November 2026", "December 2026"], [JAPAN], {"Japan Trip"})
    assert links["October 2026"] == {"X5": 0, "X6": "=X3-X4-X5"}
    assert links["November 2026"] == {"X5": "='Japan Trip'!Q4", "X6": "='October 2026'!X6+X3-X4-X5"}
    assert links["December 2026"]["X5"] == 0


def test_local_xlsx(tmp_path: Path):
    cfg = config(trips=[JAPAN])
    (path,) = write_accounts_xlsx(tmp_path, plan(build_rows(sample() + japan_sample(), cfg), cfg), cfg)
    assert path.name == "2026 Accounts.xlsx"
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["To Update", "March 2026", "July 2026", "August 2026", "November 2026", "Japan Trip", "Year End Summary"]
    ws = wb["July 2026"]
    assert [ws[f"{c}3"].value for c in "BCDEG"] == ["GRAB* RIDE", "Transport", 12.0, "Essential", "GRAB* RIDE"]
    assert [ws[f"{c}3"].value for c in "QR"] == ["SALARY", 3000.0]
    assert (ws["F3"].value, ws["S3"].value) == ("TNG eWallet", "CIMB")  # Source: the bank or wallet
    assert ws.column_dimensions["H"].hidden and ws.column_dimensions["U"].hidden
    assert wb["August 2026"]["X6"].value == "='July 2026'!X6+X3-X4-X5"
    assert wb["November 2026"]["X5"].value == "='Japan Trip'!Q4"
    trip = wb["Japan Trip"]
    assert [trip[f"{c}4"].value for c in "BD"] == ["AIRASIA BERHAD", 225.0]
    summary = wb["Year End Summary"]
    assert summary["B2"].value.endswith("+B4") and summary["B4"].value == "='Japan Trip'!Q4"


# ---- Google Sheets sync against a fake gspread --------------------------


class FakeWS:
    def __init__(self, title, id_):
        self.title, self.id = title, id_


class FakeSheet:
    url = "https://docs.google.com/spreadsheets/d/fake"

    def __init__(self, tabs=("Sheet1",)):
        self.tabs = [FakeWS(t, i) for i, t in enumerate(tabs)]
        self.values: dict[str, list[list]] = {}  # range -> rows, for values_batch_get
        self.writes: list[dict] = []
        self.requests: list[dict] = []

    def worksheets(self):
        return list(self.tabs)

    def add_worksheet(self, title, rows, cols, index=None):
        ws = FakeWS(title, len(self.tabs) + 100)
        self.tabs.insert(len(self.tabs) if index is None else index, ws)
        return ws

    def del_worksheet(self, ws):
        self.tabs.remove(ws)

    def values_batch_get(self, ranges, params=None):
        return {"valueRanges": [{"values": self.values.get(r, [])} for r in ranges]}

    def values_batch_update(self, body):
        self.writes = body["data"]

    def batch_update(self, body):
        self.requests = body["requests"]

    def values_clear(self, rng):
        self.cleared = rng


class FakeClient:
    def __init__(self):
        self.sheets: dict[str, FakeSheet] = {}

    def open(self, title):
        if title not in self.sheets:
            raise gspread.SpreadsheetNotFound
        return self.sheets[title]

    def create(self, title, folder_id=None):
        self.sheets[title] = FakeSheet()
        return self.sheets[title]


@pytest.fixture
def fake(monkeypatch):
    client = FakeClient()
    monkeypatch.setattr(gsheets, "connect", lambda g: client)
    return client


def writes(sheet):
    return {w["range"]: w["values"] for w in sheet.writes}


def titles(sheet):
    return [t.title for t in sheet.tabs]


def test_first_upload_creates_year_sheet(fake):
    cfg = config()
    (result,) = gsheets.upload(plan(build_rows(sample(), cfg), cfg), cfg)

    sheet = fake.sheets["2026 Accounts"]
    assert titles(sheet) == ["To Update", "July 2026", "August 2026", "Year End Summary"]
    assert (result.added, result.already) == (7, 0)
    w = writes(sheet)
    assert w["'July 2026'!A3:H7"][0][:5] == ["2026-07-02", "GRAB* RIDE", "Transport", 12.0, "Essential"]
    assert w["'July 2026'!P3:U3"][0][:3] == ["2026-07-01", "SALARY", 3000.0]
    assert w["'August 2026'!X6"] == [["='July 2026'!X6+X3-X4-X5"]]
    assert w["'July 2026'!J3"] == [["Bills"]]
    # Section headers are coloured.
    assert any("backgroundColor" in str(r) for r in sheet.requests)
    assert any("sortRange" in r for r in sheet.requests)


def test_rerun_appends_only_new_rows_and_keeps_edits(fake):
    cfg = config()
    rows = build_rows(sample(), cfg)
    gsheets.upload(plan(rows, cfg), cfg)
    sheet = fake.sheets["2026 Accounts"]

    # The sheet now holds the first upload, with a category picked by hand
    # and a cash row typed in by hand (no key).
    july_spend = [r for r in rows if r.date.month == 7 and r.kind == "spend"]
    sheet.values["'July 2026'!A3:H"] = [[46000, r.item, "Misc", r.amount, "", "", "", r.key] for r in july_spend] + [
        [46001, "Cash: nasi lemak", "Food", 6, "Essential", "", "", ""]
    ]
    july_earn = [r for r in rows if r.date.month == 7 and r.kind == "earn"]
    sheet.values["'July 2026'!P3:U"] = [[46000, r.item, r.amount, "", "", r.key] for r in july_earn]
    aug = [r for r in rows if r.date.month == 8]
    sheet.values["'August 2026'!A3:H"] = [[46000, r.item, "", r.amount, "", "", "", r.key] for r in aug]

    more = sample() + [tx(date(2026, 7, 20), -30.0, "GRAB* FOOD")]
    (result,) = gsheets.upload(plan(build_rows(more, cfg), cfg), cfg)

    w = writes(sheet)
    # One new row, below the 6 existing ones (rows 3-8); nothing above rewritten.
    assert w["'July 2026'!A9:H9"][0][:3] == ["2026-07-20", "GRAB* FOOD", "Transport"]
    assert not any(r.startswith("'July 2026'!A3") for r in w)
    assert (result.added, result.already) == (1, 7)
    assert titles(sheet) == ["To Update", "July 2026", "August 2026", "Year End Summary"]


def test_trip_tab_slots_in_without_moving_your_tabs(fake):
    cfg = config(trips=[JAPAN])
    # An existing year sheet with a tab you made yourself.
    fake.sheets["2026 Accounts"] = FakeSheet(["January 2026", "Chinese New Year", "November 2026", "December 2026", "Year End Summary"])

    gsheets.upload(plan(build_rows(japan_sample(), cfg), cfg), cfg)
    sheet = fake.sheets["2026 Accounts"]
    assert titles(sheet) == [
        "To Update", "January 2026", "Chinese New Year", "March 2026", "November 2026", "Japan Trip", "December 2026", "Year End Summary",
    ]
    w = writes(sheet)
    assert w["'Japan Trip'!A4:H7"][0][:2] == ["2026-08-17", "AIRASIA BERHAD"]
    assert w["'Japan Trip'!A1"] == [["Japan Trip"]]
    assert w["'November 2026'!X5"] == [["='Japan Trip'!Q4"]]
    assert w["'December 2026'!X6"] == [["='November 2026'!X6+X3-X4-X5"]]
    assert "Chinese New Year" not in " ".join(w)  # never written to


def test_trip_added_later_skips_rows_already_on_month_tab(fake):
    cfg = config()
    rows = build_rows(japan_sample(), cfg)
    gsheets.upload(plan(rows, cfg), cfg)
    sheet = fake.sheets["2026 Accounts"]
    nov = [r for r in rows if r.date.month == 11 and r.kind == "spend"]
    sheet.values["'November 2026'!A3:H"] = [[46000, r.item, "", r.amount, "", "", "", r.key] for r in nov]

    # Now the trip is configured: rows already on November aren't duplicated.
    cfg = config(trips=[JAPAN])
    gsheets.upload(plan(build_rows(japan_sample(), cfg), cfg), cfg)
    trip_rows = [v for r, vs in writes(sheet).items() if r.startswith("'Japan Trip'!A4") for v in vs]
    assert "AGODA HOTEL OSAKA" not in [v[1] for v in trip_rows]


def test_upload_without_credentials_explains_setup(tmp_path):
    cfg = Config()
    cfg.google.credentials = tmp_path / "missing.json"
    with pytest.raises(gsheets.UploadError, match="README"):
        gsheets.connect(cfg.google)
