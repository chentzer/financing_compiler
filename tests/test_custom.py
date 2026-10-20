"""Custom sheets (Sheets section) and importing trip / custom tabs from an
accounts workbook kept by hand. Made-up data only."""

import io
from datetime import datetime

import openpyxl

from fincomp import custom, webapi
from fincomp import settings as S
from fincomp.accounts import year_summary_cells
from fincomp.config import Config, CustomSheet, CustomTable
from fincomp.old_accounts import sheets_from_workbook, trips_from_workbook

ANG_PAO = CustomSheet("Chinese New Year", 2026, "earnings", [
    CustomTable("Ang Pao (RM)", "RM", "Person"), CustomTable("Ang Pao (SGD)", "SGD", "Person")])


def test_layout_side_by_side_with_totals():
    cells = custom.layout(ANG_PAO)
    assert [cells[f"{c}2"] for c in "ABCDE"] == ["Date", "Person", "Amount (RM)", "Remarks", "Key"]
    assert [cells[f"{c}2"] for c in "GHIJK"] == ["Date", "Person", "Amount (SGD)", "Remarks", "Key"]
    assert cells["A1"] == "Ang Pao (RM)" and cells["C1"] == "=SUM($C$3:$C)" and cells["I1"] == "=SUM($I$3:$I)"
    assert custom.data_ranges(ANG_PAO) == ["'Chinese New Year'!A3:E", "'Chinese New Year'!G3:K"]


def test_rows_add_edit_remove():
    values = {"'Chinese New Year'!A3:E": [[46070, "Auntie Mei", 50, "", "m-1"]], "'Chinese New Year'!G3:K": []}
    (rm, sgd) = custom.rows(ANG_PAO, values)
    assert (rm["total"], rm["rows"][0]["date"], rm["rows"][0]["item"]) == (50, "2026-02-17", "Auntie Mei")
    assert sgd["rows"] == [] and sgd["currency"] == "SGD"

    (w,) = custom.append(ANG_PAO, 0, values["'Chinese New Year'!A3:E"], {"date": "2026-02-18", "item": "Bob", "amount": 88})
    assert w["range"] == "'Chinese New Year'!A4:E4" and w["values"][0][:3] == ["2026-02-18", "Bob", 88.0]
    assert w["values"][0][4].startswith("m-")
    assert custom.edit(ANG_PAO, values, "m-1", {"amount": 60, "remarks": "Penang"}) == [
        {"range": "'Chinese New Year'!C3", "values": [[60.0]]}, {"range": "'Chinese New Year'!D3", "values": [["Penang"]]}]
    (rm_req,) = custom.remove(ANG_PAO, 9, values, "m-1")
    assert rm_req["deleteRange"]["range"] == {"sheetId": 9, "startRowIndex": 2, "endRowIndex": 3,
                                              "startColumnIndex": 0, "endColumnIndex": 5}


def test_rm_tables_count_in_the_year_end_summary():
    earn, spend = custom.summary_refs([ANG_PAO], {"Chinese New Year"}, 2026)
    assert earn == ["'Chinese New Year'!C1"] and spend == []  # SGD table isn't added
    cells = year_summary_cells(["January 2026"], [], earn, spend)
    assert cells["B1"] == "='January 2026'!X3+'Chinese New Year'!C1"
    assert custom.summary_refs([ANG_PAO], {"Chinese New Year"}, 2025) == ([], [])


def test_settings_round_trip():
    cfg = Config(sheets=[ANG_PAO])
    tabs = S.tabs_from_config(cfg)
    assert tabs[S.SHEETS][1] == ["Chinese New Year", 2026, "earnings",
                                 "Ang Pao (RM) / RM / Person ; Ang Pao (SGD) / SGD / Person"]
    assert S.config_from_tabs(tabs).sheets == [ANG_PAO]
    webapi.use_settings(tabs)
    assert webapi.sheets_list()[0]["tables"][1] == {"name": "Ang Pao (SGD)", "currency": "SGD", "label": "Person"}


def workbook() -> bytes:
    wb = openpyxl.Workbook()
    m = wb.active
    m.title = "November 2023"
    for col, h in zip("ABCDEF", ["Date", "Item", "Category", "Amount (RM)", "Importance", "Remarks"]):
        m[f"{col}2"] = h
    m.append([datetime(2023, 11, 2), "Lunch", "Food", 12.5, "Essential"])
    cny = wb.create_sheet("Chinese New Year")
    cny["A1"], cny["E1"] = "Ang Pao", "Ang Pao"
    for col, h in zip("ABCDEFGH", ["Date", "Person", "Amount (RM)", "Remarks", "Date", "Person", "Amount (SGD)", "Remarks"]):
        cny[f"{col}2"] = h
    cny.append([datetime(2023, 1, 22), "Auntie Mei", 50, "", datetime(2023, 1, 30), "SG cousin", 10])
    cny.append([datetime(2023, 1, 22), "Uncle Bob", 100])
    cny["C10"] = "=SUM(C3:C9)"  # the old total row: not a row
    trip = wb.create_sheet("Japan Trip")
    trip["A1"] = "Japan Trip"
    heads = {"A": ["Date", "Item", "Category", "Amount (RM)", "Importance", "Remarks"],
             "H": ["Date", "Item", "Category", "Amount (Yen)", "Importance", "Remarks"],
             "Q": ["Date", "Item", "Amount (Yen)", "Remarks"]}
    for start, names in heads.items():
        for i, h in enumerate(names):
            trip.cell(3, ord(start) - 64 + i, h)
    trip["A4"], trip["B4"], trip["D4"] = datetime(2023, 10, 1), "Flight", 1200
    trip["A5"], trip["B5"], trip["D5"] = datetime(2023, 10, 28), "Money Exchange", 320
    for r, (d, item, yen) in enumerate([(datetime(2023, 11, 9), "Ramen", 1200), (datetime(2023, 11, 12), "Train", 800)], 4):
        trip[f"H{r}"], trip[f"I{r}"], trip[f"J{r}"], trip[f"K{r}"] = d, item, "Food", yen
    trip["Q4"], trip["R4"], trip["S4"] = datetime(2023, 10, 28), "Money Exchange", 10000
    trip["R5"], trip["S5"] = "Bob paid back", 500  # no date: still money received
    trip["O13"], trip["P13"] = "Total Spent in Yen", 2000.0
    trip["O14"], trip["P14"] = "Total Spent in MYR (rate)", 64.0
    ys = wb.create_sheet("Year End Summary")
    ys["B1"] = "='November 2023'!T3+'Chinese New Year'!C10"
    ys["B2"] = "='November 2023'!T4+'Japan Trip'!P13"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_trip_tab_imported_with_its_three_blocks():
    (trip,) = trips_from_workbook(workbook())
    assert (trip["name"], trip["currency"], trip["start"], trip["end"], trip["year"]) == (
        "Japan Trip", "Yen", "2023-11-09", "2023-11-12", 2023)
    assert trip["rate"] == 0.032  # the old sheet's own rate
    assert [(r["item"], r["amount"]) for r in trip["rm"]] == [("Flight", 1200.0), ("Money Exchange", 320.0)]
    assert [r["item"] for r in trip["cash"]] == ["Ramen", "Train"]
    assert [(r["date"], r["item"], r["amount"]) for r in trip["received"]] == [
        ("2023-10-28", "Money Exchange", 10000.0), ("", "Bob paid back", 500.0)]


def test_trip_extras_not_twice():
    (trip,) = trips_from_workbook(workbook())
    first = webapi.trip_extras("Japan Trip", [], [], trip["cash"], trip["received"])
    assert [d["range"] for d in first] == ["'Japan Trip'!J4:O5", "'Japan Trip'!T4:W5"]
    now_cash = [[45239, "Ramen", "Food", 1200, "", ""], [45242, "Train", "Food", 800, "", ""]]
    now_received = [[45227, "Money Exchange", 10000, ""], ["", "Bob paid back", 500, ""]]
    assert webapi.trip_extras("Japan Trip", now_cash, now_received, trip["cash"], trip["received"]) == []


def test_custom_tab_imported_counted_as_earnings():
    (sheet,) = sheets_from_workbook(workbook(), {"Japan Trip"})
    assert (sheet["name"], sheet["year"], sheet["counts"]) == ("Chinese New Year", 2023, "earnings")
    rm, sgd = sheet["tables"]
    assert (rm["name"], rm["currency"], rm["label"]) == ("Ang Pao (RM)", "RM", "Person")
    assert [(r["item"], r["amount"]) for r in rm["rows"]] == [("Auntie Mei", 50.0), ("Uncle Bob", 100.0)]
    assert (sgd["currency"], [r["item"] for r in sgd["rows"]]) == ("SGD", ["SG cousin"])
    s = CustomSheet("Chinese New Year", 2023, "earnings", [CustomTable(t["name"], t["currency"], t["label"]) for t in sheet["tables"]])
    assert custom.append_many(s, 0, [[44948, "Auntie Mei", 50, "", rm["rows"][0]["key"]]], rm["rows"])[0]["values"][0][1] == "Uncle Bob"
