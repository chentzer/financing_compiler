"""Local copy of the accounts in the template layout, one workbook per year.

Rebuilt from scratch on every run; your edits live in the Google Sheet."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .accounts import (
    CAT_COLS, COLOURS, EARN_COLS, EARN_KEY, IMP_COLS, PAID_COLS, PAID_KEY, SPEND_COLS, SPEND_KEY, SUMMARY_COLS,
    TO_UPDATE, TRIP_SUMMARY_COLS, XLSX_LAST_ROW, YEAR_SUMMARY, Tab, currency_cols, month_layout, month_links,
    rate_cells, received_cols, to_update_table, trip_layout, year_summary_cells,
)
from .config import Config

MONTH_WIDTHS = {"A": 12, "B": 34, "C": 16, "D": 12, "E": 14, "F": 14, "G": 30, "I": 3, "J": 18, "K": 14, "L": 3,
                "M": 16, "N": 14, "O": 3, "P": 12, "Q": 30, "R": 12, "S": 14, "T": 30, "V": 3, "W": 12, "X": 14}
TRIP_WIDTHS = {"A": 12, "B": 30, "C": 14, "D": 12, "E": 14, "F": 14, "G": 24, "I": 3, "J": 12, "K": 22, "L": 12,
               "M": 20, "O": 3, "P": 26, "Q": 14, "R": 3, "S": 12, "T": 28, "U": 14, "V": 14, "W": 14, "X": 18}


def _fill(ws, cells: str, colour: str) -> None:
    c = COLOURS[colour]
    rgb = "".join(f"{round(c[k] * 255):02X}" for k in ("red", "green", "blue"))
    for row in ws[cells]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=rgb)


def write_accounts_xlsx(out_dir: Path, tabs_by_year: dict[int, list[Tab]], config: Config) -> list[Path]:
    paths = []
    for year, tabs in tabs_by_year.items():
        wb = Workbook()
        wb.remove(wb.active)
        months = [t.title for t in tabs if t.kind == "month"]
        trips = [t.title for t in tabs if t.kind == "trip"]
        links = month_links(months, config.trips, set(trips))
        _fill_to_update(wb.create_sheet(TO_UPDATE), tabs)
        for tab in tabs:
            ws = wb.create_sheet(tab.title)
            if tab.kind == "month":
                _fill_month(ws, tab, config, links[tab.title])
            else:
                _fill_trip(ws, tab, config)
        summary = wb.create_sheet(YEAR_SUMMARY)
        for cell, value in year_summary_cells(months, trips).items():
            summary[cell] = value
        summary.column_dimensions["A"].width = 18
        summary.column_dimensions["B"].width = 14

        path = out_dir / (config.google.spreadsheet.replace("{year}", str(year)) + ".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        paths.append(path)
    return paths


def _fill_to_update(ws, tabs: list[Tab]) -> None:
    for row in to_update_table(tabs):
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[1].number_format = "+#,##0.00;-#,##0.00"
    for col, w in {"A": 12, "B": 12, "C": 26, "D": 56, "E": 26, "F": 26, "G": 16}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "A2"


def _rows(ws, tab: Tab, kind: str, cols: str, date_col: str, amount_col: str) -> None:
    for i, r in enumerate((r for r in tab.rows if r.kind == kind), start=tab.first_row):
        for col, v in zip(cols, r.values()):
            ws[f"{col}{i}"] = v
        ws[f"{date_col}{i}"].number_format = "dd/mm/yyyy"
        ws[f"{amount_col}{i}"].number_format = "#,##0.00"


def _dropdowns(ws, config: Config, first_row: int, pairs: list[tuple[str, list[str]]]) -> None:
    for col, values in pairs:
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.add(f"{col}{first_row}:{col}{XLSX_LAST_ROW}")
        ws.add_data_validation(dv)


def _fill_month(ws, tab: Tab, config: Config, links: dict) -> None:
    for cell, value in {**month_layout(config, open_ended=False), **links}.items():
        ws[cell] = value
    for cell in ("A1", f"{EARN_COLS[0]}1"):
        ws[cell].font = Font(bold=True, size=12)
    _rows(ws, tab, "spend", SPEND_COLS + SPEND_KEY, "A", "D")
    _rows(ws, tab, "earn", EARN_COLS + EARN_KEY, EARN_COLS[0], EARN_COLS[2])
    for col in (CAT_COLS[1], IMP_COLS[1], SUMMARY_COLS[1]):
        for row in range(tab.first_row, tab.first_row + len(config.categories) + 2):
            ws[f"{col}{row}"].number_format = "#,##0.00"
    _dropdowns(ws, config, tab.first_row, [("C", config.categories), ("E", config.importance)])
    for col, w in MONTH_WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[SPEND_KEY].hidden = True
    ws.column_dimensions[EARN_KEY].hidden = True
    _fill(ws, f"A1:{SPEND_COLS[-1]}2", "spend")
    _fill(ws, f"{EARN_COLS[0]}1:{EARN_COLS[-1]}2", "earn")
    for cols in (CAT_COLS, IMP_COLS, SUMMARY_COLS):
        _fill(ws, f"{cols[0]}2:{cols[1]}2", "summary")
    _fill(ws, f"{SUMMARY_COLS[0]}6:{SUMMARY_COLS[1]}6", "final")
    ws.freeze_panes = "A3"


def _fill_trip(ws, tab: Tab, config: Config) -> None:
    for cell, value in trip_layout(tab.trip, config, open_ended=False).items():
        ws[cell] = value
    ws["A1"].font = Font(bold=True, size=14)
    _rows(ws, tab, "spend", SPEND_COLS + SPEND_KEY, "A", "D")
    L, V = TRIP_SUMMARY_COLS
    currencies = tab.trip.currencies()
    for row in range(4, 9 + 4 * len(currencies)):
        ws[f"{V}{row}"].number_format = "#,##0.00"
    for cell in rate_cells(tab.trip):
        ws[cell].number_format = "0.000000"
    pairs = [("C", config.categories), ("E", config.importance)]
    for i, _ in enumerate(currencies):
        c = currency_cols(i)
        pairs += [(c[2], config.categories), (c[4], config.importance)]
        _fill(ws, f"{c[0]}2:{c[-1]}3", "cash")
    if currencies:
        rc = received_cols(len(currencies))
        _fill(ws, f"{rc[0]}2:{rc[-1]}3", "summary")
    _fill(ws, f"A2:{SPEND_COLS[-1]}3", "spend")
    _fill(ws, f"{PAID_COLS[0]}2:{PAID_COLS[-1]}3", "earn")
    _fill(ws, f"{L}3:{V}3", "summary")
    _fill(ws, f"{L}4:{V}4", "final")
    ws.column_dimensions[PAID_KEY].hidden = True
    _dropdowns(ws, config, tab.first_row, pairs)
    for col, w in TRIP_WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[SPEND_KEY].hidden = True
    ws.freeze_panes = "A4"
