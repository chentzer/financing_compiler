from __future__ import annotations

from dataclasses import asdict
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from .combine import TransferMatch
from .models import Transaction

HEADERS = {
    "date": "Date",
    "source": "Source",
    "account": "Account",
    "type": "Type",
    "description": "Description",
    "amount": "Amount (RM)",
    "reference": "Reference",
    "internal": "Internal",
    "note": "Note",
    "file": "File",
}


def to_frame(transactions: list[Transaction]) -> pd.DataFrame:
    df = pd.DataFrame([asdict(t) for t in transactions], columns=Transaction.columns())
    return df[list(HEADERS)].rename(columns=HEADERS)


def monthly_summary(df: pd.DataFrame) -> pd.DataFrame:
    real = df[~df["Internal"]].copy()
    if real.empty:
        return pd.DataFrame(columns=["Month", "Source", "Money In", "Money Out", "Net"])
    real["Month"] = pd.to_datetime(real["Date"]).dt.strftime("%Y-%m")
    amount = real["Amount (RM)"]
    real["Money In"] = amount.where(amount > 0, 0.0)
    real["Money Out"] = amount.where(amount < 0, 0.0)
    per_source = real.groupby(["Month", "Source"], as_index=False)[["Money In", "Money Out"]].sum()
    totals = real.groupby("Month", as_index=False)[["Money In", "Money Out"]].sum().assign(Source="All")
    out = pd.concat([per_source, totals], ignore_index=True)
    out["Net"] = out["Money In"] + out["Money Out"]
    out = out.sort_values(["Month", "Source"], key=lambda s: s.where(s != "All", "￿"))
    return out[["Month", "Source", "Money In", "Money Out", "Net"]].round(2)


def transfers_frame(matches: list[TransferMatch]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Out Date": m.out.date,
                "From": f"{m.out.source} {m.out.account}",
                "Out Description": m.out.description,
                "In Date": m.into.date,
                "To": f"{m.into.source} {m.into.account}",
                "In Description": m.into.description,
                "Amount (RM)": m.into.amount,
            }
            for m in matches
        ],
        columns=["Out Date", "From", "Out Description", "In Date", "To", "In Description", "Amount (RM)"],
    )


def write_excel(path: Path, transactions: list[Transaction], matches: list[TransferMatch]) -> None:
    df = to_frame(transactions)
    real = df[~df["Internal"]]
    sheets = {
        "All Transactions": df,
        "Money In": real[real["Amount (RM)"] > 0],
        "Money Out": real[real["Amount (RM)"] < 0],
        "Monthly Summary": monthly_summary(df),
        "Transfers": transfers_frame(matches),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            _format(writer.sheets[name], frame)


def write_csv(path: Path, transactions: list[Transaction]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    to_frame(transactions).to_csv(path, index=False, encoding="utf-8-sig")


def _format(ws, frame: pd.DataFrame) -> None:
    ws.freeze_panes = "A2"
    for i, col in enumerate(frame.columns, start=1):
        values = [str(col)] + [str(v) for v in frame[col].head(500)]
        ws.column_dimensions[get_column_letter(i)].width = min(max(map(len, values)) + 2, 60)
        if "RM" in str(col) or col in ("Money In", "Money Out", "Net"):
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = "#,##0.00"
        elif "Date" in str(col):
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = "yyyy-mm-dd"
